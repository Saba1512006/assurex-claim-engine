"""Build the submission documents in the faculty's folder layout (submission/Documents/).

    python reports/build_submission_documents.py

Writes Project_Report.pdf, Python_Model_Report.pdf, Teachable_Machine_Report.pdf, Model_Comparison_Report.pdf,
Test_Cases.xlsx, Test_Results.pdf, Installation_Guide.pdf, User_Guide.pdf, Technical_Blog_Link.txt,
Team_Contributions.md and Deployment_URL.txt. Every document is generated from the project's own sources
(see reports/build_documentation.py), so re-running this after a change keeps them in step.
`scripts/make_submission.ps1` then assembles Source_Code/, Documents/ and Videos/ on Windows.
"""
from __future__ import annotations

import csv
import html
import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_documentation as bd  # noqa: E402

ROOT, DOC = bd.ROOT, bd.DOC
OUT = ROOT / "submission" / "Documents"
LIVE = "https://assurex.pythonanywhere.com"
BLOG = "https://medium.com/@sabarajput672/building-assurex-two-models-one-rulebook-and-why-our-first-100-was-a-bug-bb8b9858b11c"
VIDEO = f"{bd.REPO}/blob/demo-video/AssureX_Demo.mp4"
ACCOUNTS = [("Customer", "customer@assurex.local", "CustomerPass123!"),
            ("Service-center staff", "staff@assurex.local", "StaffPass123!"),
            ("Claim reviewer", "reviewer@assurex.local", "ReviewerPass123!"),
            ("Administrator", "admin@assurex.local", "AdminPass123!")]


# ---------------------------------------------------------------- test results
def test_results() -> list[tuple[str, str, str]]:
    """(file, test, status) for every test in the latest recorded run."""
    out = []
    for ln in bd.read("reports/test_results.txt").splitlines():
        m = re.match(r"\s+(PASSED|FAILED|ERROR|SKIPPED)\s+(tests/\S+?)::(.+?)\s*$", ln)
        if m:
            out.append((m.group(2), m.group(3), m.group(1)))
    return out


def test_case_rows() -> list[list]:
    """One row per test function named in documentation/TEST_CASES.md, with its recorded result."""
    results = test_results()
    rows, n = [], 0
    for ln in bd.read("documentation/TEST_CASES.md").splitlines():
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if not ln.startswith("| **") or len(cells) < 3:
            continue
        category, checked = cells[0].strip("*"), cells[1]
        for group in re.findall(r"`([^`]+)`", cells[2]):
            if "::" not in group:
                continue
            module, funcs = (x.strip() for x in group.split("::", 1))
            path = module if module.endswith(".py") else f"tests/{module}.py"
            for func in (f.strip() for f in funcs.split(",")):
                if not func:
                    continue
                pattern = re.escape(func).replace(r"\*", ".*")
                hits = [r for r in results if r[0] == path
                        and re.fullmatch(rf"(test_[^\[]*\[)?{pattern}(\[.*\])?\]?", r[1])]
                passed = sum(1 for h in hits if h[2] == "PASSED")
                n += 1
                status = "Pass" if hits and passed == len(hits) else ("Fail" if hits else "Not run separately")
                actual = (f"{passed} of {len(hits)} passed" if hits else "covered by the named file")
                rows.append([f"TC-{n:03d}", category, checked, f"{path} :: {func}",
                             "Test passes (behaviour as described)", actual, status])
    return rows


def write_test_cases_xlsx(path: Path) -> None:
    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="1D5BD8"), Font(bold=True, color="FFFFFF")

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    results = test_results()
    passed = sum(1 for r in results if r[2] == "PASSED")
    ws = wb.active
    ws.title = "Summary"
    summary = [("Project", "AssureX Claim Engine — Group NN_DevStorm"),
               ("Test command", "python -m pytest -q  (unit, integration, security, model and rule tests)"),
               ("Browser tests", "python -m pytest -m e2e tests/e2e  (Playwright, accessibility checks)"),
               ("Recorded tests", f"{len(results)} — {passed} passed, {len(results) - passed} not passed"),
               ("Source", "reports/test_results.txt; test code in tests/")]
    for line in bd.read("reports/test_results.txt").splitlines()[:8]:
        if ":" in line and line.split(":")[0].strip() in ("Date", "Python", "Result", "Browser"):
            k, v = line.split(":", 1)
            summary.append((k.strip(), v.strip()))
    sheet(ws, ["Item", "Value"], summary, [22, 100])

    sheet(wb.create_sheet("Test cases"),
          ["ID", "SRS category", "What is checked", "Automated test", "Expected result", "Actual result", "Status"],
          test_case_rows(), [9, 22, 60, 55, 26, 22, 12])

    demo = []
    for ln in bd.read("sample_claims/README.md").splitlines():
        m = re.match(r"\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*(.+?)\s*\|\s*(D\d+)\s*\|", ln)
        if m:
            f = next((p.name for p in (ROOT / "sample_claims").glob(f"{int(m.group(1)):02d}_*.json")), "")
            hit = [r for r in results if r[1] == f"test_demonstration_case[{Path(f).stem}]"]
            status = "Pass" if hit and hit[0][2] == "PASSED" else "—"
            demo.append([f"DC-{int(m.group(1)):02d}", m.group(2), f"sample_claims/{f}", m.group(3), m.group(4),
                         f"{m.group(3)} ({m.group(4)})" if status == "Pass" else "", status])
    sheet(wb.create_sheet("Demonstration cases"),
          ["ID", "Case", "Input", "Expected recommendation", "Decision row", "Actual result", "Status"],
          demo, [8, 42, 42, 26, 12, 34, 10])

    readiness = [(ln.strip()[6:],) for ln in bd.read("documentation/TEST_CASES.md").splitlines() if ln.strip().startswith("- [x]")]
    sheet(wb.create_sheet("Hidden-test readiness"), ["Checked"], readiness, [120])

    sheet(wb.create_sheet("All automated tests"), ["Test file", "Test", "Result"],
          [list(r) for r in results], [45, 90, 10])
    wb.save(path)
    print("written", path.relative_to(ROOT))


def test_results_chapter() -> str:
    results = test_results()
    per_file = {}
    for f, _, s in results:
        per_file.setdefault(f, [0, 0])[0 if s == "PASSED" else 1] += 1
    head = [ln for ln in bd.read("reports/test_results.txt").splitlines()[:8] if ln.strip()]
    return ("<h2>Summary</h2><pre>" + html.escape("\n".join(head)) + "</pre>"
            + "<h2>Results by test file</h2>"
            + bd.table(["Test file", "Passed", "Not passed"], [(k, v[0], v[1]) for k, v in sorted(per_file.items())], "small")
            + "<h2>Every test</h2>" + bd.table(["Test file", "Test", "Result"], results, "tiny"))


def python_model_chapter() -> str:
    card = json.loads(bd.read("model/python_model/model_card_v2.json"))
    parts = [bd.md(bd.read("documentation/PYTHON_MODEL_EVIDENCE.md"), DOC)]
    imp = card.get("permutation_importance_val", {})
    if imp:
        parts.append("<h2>Feature importance (permutation, validation split)</h2>"
                     + bd.table(["Feature", "Importance (drop in macro-F1)"], list(imp.items())))
    with open(ROOT / "model/python_model/sample_test_predictions.csv", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    parts.append("<h2>Sample test predictions (first 25 of the 225 test claims)</h2>"
                 "<p>Full list: <code>model/python_model/sample_test_predictions.csv</code>.</p>"
                 + bd.table(rows[0], rows[1:26], "small"))
    files = [(p.name, f"{p.stat().st_size / 1024:.0f} KB") for p in sorted((ROOT / "model/python_model").iterdir()) if p.is_file() and not p.name.startswith(".")]
    parts.append("<h2>Saved model files (<code>model/python_model/</code>)</h2>" + bd.table(["File", "Size"], files))
    return "\n".join(parts)


def gtm_samples() -> str:
    parts = ["<h2>Class samples</h2><p>Three training cards per class from "
             "<code>data/summary_cards/train/</code> (700 per class were uploaded to Teachable Machine).</p>"]
    for cls, label in (("valid", "Valid Claim"), ("invalid", "Invalid Claim"), ("manual_review", "Manual Review")):
        cards = sorted((ROOT / "data/summary_cards/train" / cls).glob("*.jpg"))[:3]
        parts.append(f"<h3>{label}</h3><div class=\"cards\">"
                     + "".join(bd.figure(c, c.name, 600, box_w=200) for c in cards) + "</div>")
    return "\n".join(parts)


# ---------------------------------------------------------------- documents
def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    report = bd.read("documentation/PROJECT_REPORT.md")
    report = re.sub(r"\A# .*?(?=^## 1\. )", "", report, flags=re.S | re.M)
    report_html = bd.md(report, DOC).replace("<h2>13. Data dictionary</h2>", "<h2>13. Data dictionary</h2>" + bd.data_dictionary_html())
    docs = [
        ("Project_Report.pdf", "Project Report", [
            ("Project report", report_html), ("Dataset", bd.dataset_chapter()),
            ("Warranty policy files", bd.policies_chapter()),
            ("Security and access-control design", bd.md(bd.read("documentation/RBAC_DESIGN.md"), DOC)),
            ("Demonstration video", bd.md(bd.read("documentation/DEMO_VOICEOVER.md"), DOC)),
            ("AI tool usage declaration", bd.md(bd.read("AI_USAGE.md"), ROOT)),
            ("Development log", bd.md(bd.read("documentation/DEVELOPMENT_LOG.md"), DOC))],
         "The project report required by SRS §1.10 item 1 (all 35 sections, with the data flow, use case, activity, "
         "sequence and decision-flow diagrams), followed by the dataset, the warranty policy files, the security design, "
         "the demonstration video map, the AI usage declaration and the development log."),
        ("Python_Model_Report.pdf", "Python Classification Model Report",
         [("Python classification model evidence", python_model_chapter())],
         "Complete evidence for the Python classification model (SRS §1.10 item 4): data, pre-processing, the three "
         "algorithms compared, hyper-parameters, cross-validation, validation and test results, confusion matrix, "
         "per-class metrics, feature importance, saved files and sample predictions."),
        ("Teachable_Machine_Report.pdf", "Google Teachable Machine Report",
         [("Google Teachable Machine evidence", bd.gtm_chapter() + gtm_samples())],
         "Evidence for the Google Teachable Machine model (SRS §1.10 item 5): classes and samples, training "
         "configuration and observations, incorrectly classified samples, retraining, the exported model and label "
         "file, and screenshots of training, testing and evaluation."),
        ("Model_Comparison_Report.pdf", "Model Prediction and Confidence Comparison Report",
         [("Model prediction and confidence comparison", bd.comparison_chapter(limit=None))],
         "The model comparison report (SRS §1.10 item 6) on 225 unseen test claims, each with a structured record "
         "for the Python model and its Claim Summary Card for Google Teachable Machine."),
        ("Test_Results.pdf", "Test Results",
         [("Test cases", bd.md(bd.read("documentation/TEST_CASES.md"), DOC)), ("Test results", test_results_chapter())],
         "The test cases by SRS category (§1.10 item 8) and the results of the latest automated run. The same test "
         "cases are listed one per row in Test_Cases.xlsx."),
        ("Installation_Guide.pdf", "Installation Guide",
         [("Installation instructions", bd.md(bd.read("documentation/INSTALLATION.md"), DOC))],
         "Installation instructions (SRS §1.10 item 9): prerequisites, supported systems, Python version, virtual "
         "environment, dependencies, database set-up and seeding, model placement, OCR, environment variables, run "
         "and test commands, the default administrator and troubleshooting."),
        ("User_Guide.pdf", "User Guide",
         [("Execution instructions and user guide", bd.guide_chapter()), ("Deployed application", bd.deployment_chapter())],
         "How to use the application (SRS §1.10 item 10), from registration to exporting a claim report, with "
         "screens of the working application for every role, and how to test the deployed version."),
    ]
    for name, title, chs, about in docs:
        bd.render(bd.build_html(chs, title, about), OUT / name, title)

    write_test_cases_xlsx(OUT / "Test_Cases.xlsx")
    shutil.copy(DOC / "TEAM_CONTRIBUTIONS.md", OUT / "Team_Contributions.md")
    (OUT / "Technical_Blog_Link.txt").write_text(
        "AssureX Claim Engine — Technical Blog (Group NN_DevStorm)\n\n"
        f"Published blog (Medium):\n{BLOG}\n\n"
        "The same text is served by the application at /blog and kept in Source_Code/documentation/TECHNICAL_BLOG.md.\n",
        encoding="utf-8")
    (OUT / "Deployment_URL.txt").write_text(
        "AssureX Claim Engine — Deployment (Group NN_DevStorm)\n\n"
        f"Public application URL:\n{LIVE}\n\n"
        f"Source code (public GitHub repository):\n{bd.REPO}\n\n"
        f"Demonstration video (download):\n{VIDEO}\n\n"
        "Evaluator and administrator logins (demo data only):\n"
        + "".join(f"  {r:<22} {e:<26} {p}\n" for r, e, p in ACCOUNTS)
        + "\nSample claim records: Source_Code/sample_claims/ (the 11 SRS demonstration cases); the seeded database\n"
          "contains the same eleven cases as real claims.\n\n"
          "Testing the application: see User_Guide.pdf, chapter 2 (Deployed application).\n\n"
          "Note: the free PythonAnywhere account runs Python 3.10 with a 512 MB disk quota, so both models can show as\n"
          "'Unavailable' there and claims then go to manual review. The full pipeline runs locally (Installation_Guide.pdf)\n"
          "and is shown working in Videos/Demonstration.mp4.\n",
        encoding="utf-8")
    for f in ("Team_Contributions.md", "Technical_Blog_Link.txt", "Deployment_URL.txt"):
        print("written", (OUT / f).relative_to(ROOT))


if __name__ == "__main__":
    build()
